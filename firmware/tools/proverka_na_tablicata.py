# -*- coding: utf-8 -*-
"""Проверка на формулите в работната книга с библиотеката formulas (без LibreOffice)."""
import datetime as dt
import math
import os
import shutil
import sys
import tempfile

import formulas
import openpyxl

SRC = sys.argv[1]
tmp = tempfile.mkdtemp()


def compute(path):
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    out = {}
    for k, v in sol.items():
        key = str(k)
        if "!" not in key:
            continue
        sheet, ref = key.rsplit("!", 1)
        sheet = sheet.split("]", 1)[1].rstrip("'")
        try:
            val = v.value[0][0]
        except Exception:                      # noqa: BLE001
            val = v
        out[(sheet.upper(), ref.upper())] = val
    return out


def formula_cells(path):
    wb = openpyxl.load_workbook(path)
    res = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    res.append((ws.title, c.coordinate, c.value))
    return res


def is_error(v):
    s = str(v)
    return s.startswith("#") or "XlError" in type(v).__name__


fails = 0


def check(cond, msg):
    global fails
    if not cond:
        fails += 1
        print("FAIL:", msg)


# ---- 1. Празна книга: никоя формула не връща грешка -------------------------
cells = formula_cells(SRC)
vals = compute(SRC)
errs = []
for sh, ref, f in cells:
    v = vals.get((sh.upper(), ref))
    if v is None:
        errs.append((sh, ref, f, "няма стойност"))
    elif is_error(v):
        errs.append((sh, ref, f, v))
print(f"Празна книга: {len(cells)} формули, грешки: {len(errs)}")
for e in errs[:20]:
    print("   ", e)
check(not errs, "грешки в празната книга")

# ---- 2. Книга с примерни входни данни ---------------------------------------
test = os.path.join(tmp, "test.xlsx")
shutil.copy(SRC, test)
wb = openpyxl.load_workbook(test)
w = wb["Е3 VCC"]
w["J7"] = 100; w["J8"] = 47
uvin = [5.00, 4.98, 4.96, 4.94, 4.92]; umod = [4.95, 4.90, 4.85, 4.80, 4.75]
vui = [5.01, 4.97, 4.95, 4.95, 4.91]; a0 = [470, 470, 470.5, 471, 471]
for n in range(5):
    r = 14 + n
    w[f"C{r}"] = uvin[n]; w[f"D{r}"] = umod[n]; w[f"E{r}"] = 1.45; w[f"F{r}"] = vui[n]; w[f"G{r}"] = a0[n]
w["C19"] = 5.02; w["F19"] = 5.03; w["G19"] = 471   # серия 2, 0 релета
w = wb["Е5 Ток"]
for col, v in zip("BCDEF", [10, 85, 160, 235, 310]):
    w[f"{col}7"] = v
w["B11"] = 4.75
w = wb["Е6 Бутони"]
w["C8"] = 50; w["D8"] = 120; w["E8"] = 5
w = wb["Е7 Влага"]
for r in range(12, 17):
    w[f"B{r}"] = "Да"
w["C12"] = "Не"
w = wb["Е8 PING"]
w["C10"] = 100; w["D10"] = 98; w["E10"] = 2; w["F10"] = 5000; w["G10"] = 7000; w["H10"] = 30000
w["C11"] = 100; w["D11"] = 100; w["E11"] = 0; w["F11"] = 4000; w["G11"] = 9000; w["H11"] = 40000
w = wb["Е9 Надеждност"]
w["A8"] = dt.datetime(2026, 9, 25, 22, 0); w["B8"] = dt.datetime(2026, 9, 26, 7, 30)
w["C8"] = 17000; w["D8"] = 10; w["E8"] = 17000; w["F8"] = 5; w["G8"] = 0
w = wb["Е10 Обхват"]
w["D8"] = 95; w["E8"] = 5
w = wb["Е12 Включване"]
w["B7"] = 10; w["B8"] = 3
w = wb["Е13 Уеб"]
for i in range(20):
    w[f"B{7 + i}"] = i + 1
wb.save(test)

v = compute(test)
g = lambda sh, ref: v.get((sh.upper(), ref))


def near(a, b, eps=1e-6):
    try:
        return abs(float(a) - float(b)) <= eps
    except Exception:                          # noqa: BLE001
        return False


rb = 47 * 320 / (47 + 320)
fs = 3.2 * (100 + rb) / rb
check(near(g("Е3 VCC", "J9"), fs), f"J9 {g('Е3 VCC', 'J9')} != {fs}")
check(near(g("Е3 VCC", "J10"), fs / 10.91 - 1), "J10")
for n in range(5):
    r = 14 + n
    h = a0[n] / 1024 * 10.91
    check(near(g("Е3 VCC", f"H{r}"), h), f"H{r}")
    check(near(g("Е3 VCC", f"I{r}"), (vui[n] - uvin[n]) * 1000, 1e-6), f"I{r}")
    check(near(g("Е3 VCC", f"J{r}"), (h - uvin[n]) * 1000, 1e-6), f"J{r}")
    check(near(g("Е3 VCC", f"K{r}"), (uvin[n] - umod[n]) * 1000, 1e-6), f"K{r}")
    check(near(g("Е3 VCC", f"L{r}"), 1.45 / uvin[n]), f"L{r}")
check(g("Е3 VCC", "I20") == "", "празен ред дава празно")
# средно за 0 релета: серии 1 и 2
check(near(g("Е3 VCC", "C32"), (5.00 + 5.02) / 2), f"C32 {g('Е3 VCC', 'C32')}")
check(near(g("Е3 VCC", "I32"), ((5.01 - 5.00) + (5.03 - 5.02)) / 2 * 1000, 1e-6), "I32")
check(near(g("Е3 VCC", "D33"), 4.90), "D33 (само серия 1)")
check(g("Е3 VCC", "D32") == 4.95 or near(g("Е3 VCC", "D32"), 4.95), "D32 (D19 е празна)")
iv = [(vui[n] - uvin[n]) * 1000 for n in range(5)] + [(5.03 - 5.02) * 1000]
check(near(g("Е3 VCC", "I37"), max(abs(x) for x in iv), 1e-6), f"I37 {g('Е3 VCC', 'I37')}")

check(near(g("Е5 Ток", "B10"), 75), "Е5 B10")
check(near(g("Е5 Ток", "C8"), 75) and near(g("Е5 Ток", "F8"), 75), "Е5 прираст")
check(near(g("Е5 Ток", "B12"), 4.75 * 310 / 1000), "Е5 B12")

check(near(g("Е6 Бутони", "F8"), 0) and near(g("Е6 Бутони", "G8"), 20) and near(g("Е6 Бутони", "H8"), 2.4), "Е6")
check(g("Е6 Бутони", "G9") == "", "Е6 празен ред")

check(g("Е7 Влага", "B17") == "5 от 5", f"Е7 B17 = {g('Е7 Влага', 'B17')}")
check(g("Е7 Влага", "C17") == "0 от 1", f"Е7 C17 = {g('Е7 Влага', 'C17')}")

check(near(g("Е8 PING", "M10"), 0.98) and near(g("Е8 PING", "N10"), 7.0), "Е8 ред 10")
check(near(g("Е8 PING", "D15"), 198) and near(g("Е8 PING", "E15"), 2), "Е8 сума")
check(near(g("Е8 PING", "F15"), 4000) and near(g("Е8 PING", "G15"), 8000) and near(g("Е8 PING", "H15"), 40000), "Е8 мин/ср/макс")
check(near(g("Е8 PING", "M15"), 198 / 200), "Е8 доставени общо")

check(near(g("Е9 Надеждност", "H8"), 9.5, 1e-9), f"Е9 H8 = {g('Е9 Надеждност', 'H8')}")
check(near(g("Е9 Надеждност", "I8"), 17000 / 17010), "Е9 I8")
check(near(g("Е9 Надеждност", "J8"), 17000 / 17005), "Е9 J8")
check(near(g("Е9 Надеждност", "K8"), 9.5 * 1800, 1e-6), "Е9 K8")

check(near(g("Е10 Обхват", "H8"), 0.95), "Е10")
check(near(g("Е12 Включване", "C7"), 1.0) and near(g("Е12 Включване", "C8"), 0.3), "Е12")
check(near(g("Е13 Уеб", "B28"), 10.5) and near(g("Е13 Уеб", "B29"), 1) and near(g("Е13 Уеб", "B30"), 20), "Е13")
check(g("Е13 Уеб", "C28") == "", "Е13 празна колона")

errs2 = [(sh, ref) for sh, ref, f in formula_cells(test) if is_error(v.get((sh.upper(), ref)))]
check(not errs2, f"грешки в попълнената книга: {errs2[:10]}")
print(f"Попълнена книга: {len(formula_cells(test))} формули, грешки: {len(errs2)}")
print("РЕЗУЛТАТ:", "OK" if fails == 0 else f"{fails} неуспешни проверки")
sys.exit(1 if fails else 0)
